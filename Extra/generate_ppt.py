import collections.abc
from pptx import Presentation
import shutil

template_path = r"C:\Users\nishi\OneDrive\Documents\Desktop\CSF\SEM8\Capestone\Documentation\BTech-Endterm_PPT_Template_2025-26(1).pptx"
output_path = r"C:\Users\nishi\OneDrive\Documents\Desktop\CSF\SEM8\Capestone\W-CERT_Endterm_Presentation.pptx"

shutil.copy2(template_path, output_path)

prs = Presentation(output_path)

def replace_text_in_shape(shape, old_text, new_text):
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if old_text in run.text:
                run.text = run.text.replace(old_text, new_text)

# --- SLIDE 1 ---
slide1 = prs.slides[0]
for shape in slide1.shapes:
    if shape.has_text_frame:
        # Title
        if "Title of Project" in shape.text:
            shape.text_frame.clear()
            p = shape.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = "W-CERT: A Women-Centric Cyber Emergency Response & Threat-Analysis Framework"
            r.font.bold = True
        elif "Names of Students" in shape.text:
            shape.text_frame.clear()
            p = shape.text_frame.paragraphs[0]
            p.text = "Names of Students with ERP nos:\nNishinth Venkatesh (1032220939)\nAyush Dhore (1032220255)\nNeel Karnavat (1032220469)\nSarthak Paymode (1032220213)\n\nName of BTech Capstone Project Guide: Dr. Vinayak Prabhakar Musale"


# --- SLIDE 2 & 3 ---
# They are quite static in the template, we'll leave Slide 2 as is.
# For Slide 3, we leave it as is or add specific technologies.
slide3 = prs.slides[2]
for shape in slide3.shapes:
    if shape.has_text_frame and "Literature Survey" in shape.text:
        text_frame = shape.text_frame
        for i, p in enumerate(text_frame.paragraphs):
            if "Implementation of Project" in p.text:
                p.text = "Implementation of Project [ React.js, Python/Flask, XAI (SHAP/LIME), AES-256, SHA-256 ]"


# --- SLIDE 4: INDIVIDUAL IN DETAILS ---
# We need to create a slide for EACH student based on slide 4.
# Python-pptx doesn't support duplicating slides easily.
# But we can add new slides using the same layout.
# Slide 4 is layout index 1 (Title and Content) or so. Let's check slide 4 layout name.
slide4_layout = prs.slides[3].slide_layout

students = [
    {
        "name": "Nishinth Venkatesh",
        "module": "API Architecture, Threat Intelligence Modeling, and Cloud Deployment",
        "aim": "Engineer a stateless, RESTful API backend, heuristic Threat Intelligence Engine, and resilient cloud deployment.",
        "obj": "To orchestrate the platform's core analytical routing layer and autonomous parsing of incident telemetry.",
        "contrib": "Engineered Flask API endpoints, built Threat Intelligence Engine (MITRE ATT&CK mapping), managed decoupled integration of XAI and security modules."
    },
    {
        "name": "Ayush Hemant Dhore",
        "module": "Cryptographic Infrastructure & Secure Backend Connection",
        "aim": "Establish mathematically irrefutable digital chain of custody and secure sensitive victim PII.",
        "obj": "Implement automated SHA-256 hashing pipeline and AES-256 encryption protocols.",
        "contrib": "Took ownership of the platform's cryptographic persistence layer, developed specialized security utility scripts for hashing and encryption, and managed database communication layer."
    },
    {
        "name": "Neel Harshad Karnavat",
        "module": "XAI & ML Evidence Training (Authenticity Checker)",
        "aim": "Detect pixel-level anomalies and synthetic media manipulations using Explainable AI (XAI).",
        "obj": "Develop Deep Learning model and autonomously generate human-interpretable visual proofs.",
        "contrib": "Engineered core forensic Authenticity Checker, trained deep learning network on industry-standard datasets, integrated XAI frameworks (SHAP/LIME) for dynamic heatmaps."
    },
    {
        "name": "Sarthak Anil Paymode",
        "module": "Frontend UI/UX, Console Architecture, and React/Node.js Integration",
        "aim": "Design a responsive GUI for incident reporting and real-time forensic orchestration.",
        "obj": "Architect secure asynchronous ingestion pipeline and dynamic state management framework.",
        "contrib": "Directed development of visual layers, engineered secure authentication middleware, incident ingestion forms, centralized Analyst Dashboard, and XAI heatmap overlays."
    }
]

# Modify the existing slide 4 for the first student
slide4 = prs.slides[3]
shape_title = slide4.shapes.title
shape_title.text = f"INDIVIDUAL - IN DETAILS ({students[0]['name']})"
for shape in slide4.shapes:
    if shape.has_text_frame and "Individual Aim" in shape.text:
        tf = shape.text_frame
        tf.clear()
        tf.add_paragraph().text = f"Individual Aim: {students[0]['aim']}"
        tf.add_paragraph().text = f"Individual Objective: {students[0]['obj']}"
        tf.add_paragraph().text = f"Individual contribution ({students[0]['module']}): {students[0]['contrib']}"
        tf.add_paragraph().text = "Communication & presentation skills"
        tf.add_paragraph().text = "Programming & ethics"

# Add slides for the rest of the students
for student in students[1:]:
    new_slide = prs.slides.add_slide(slide4_layout)
    new_slide.shapes.title.text = f"INDIVIDUAL - IN DETAILS ({student['name']})"
    # The layout has a body placeholder usually at index 1
    body_shape = None
    for shape in new_slide.placeholders:
        if shape.placeholder_format.idx == 1:
            body_shape = shape
            break
    if body_shape:
        tf = body_shape.text_frame
        tf.text = f"Individual Aim: {student['aim']}"
        tf.add_paragraph().text = f"Individual Objective: {student['obj']}"
        tf.add_paragraph().text = f"Individual contribution ({student['module']}): {student['contrib']}"
        tf.add_paragraph().text = "Communication & presentation skills"
        tf.add_paragraph().text = "Programming & ethics"

# Move the publications and references slides to the end (since we added slides, they are technically before the new slides unless we reorder.
# python-pptx doesn't have an easy slide reorder API. We just added slides to the end.
# Actually, slide.add_slide adds to the end! So Publication and References are now slide 5 and 6, and the new students are 7, 8, 9.
# We should move them. Let's recreate Publication and References slides at the end, and delete the original ones.
# Or better yet, just extract the XML logic to move them, or simply use the layout to recreate them and delete the old ones.

slide5_layout = prs.slides[4].slide_layout
slide6_layout = prs.slides[5].slide_layout

# Recreate Publication
pub_slide = prs.slides.add_slide(slide5_layout)
pub_slide.shapes.title.text = "Publication details"
for shape in pub_slide.placeholders:
    if shape.placeholder_format.idx == 1:
        shape.text = "Details on Literature Review Paper:\n[Insert Paper Details Here]\n\nDetails of Implementation Paper:\n[Insert Paper Details Here]"

# Recreate References
ref_slide = prs.slides.add_slide(slide6_layout)
ref_slide.shapes.title.text = "References"
for shape in ref_slide.placeholders:
    if shape.placeholder_format.idx == 1:
        shape.text = "[1] Z. Khalid, F. Iqbal, and B. C. M. Fung, 'Towards a unified XAI-based framework for digital forensic investigations,' 2024.\n[2] N. Bharati et al., 'Explainable deepfake detection,' 2025.\n[3] F. Khalid et al., 'DFP-Net: An explainable and trustworthy framework,' 2023."

# To delete old slides 5 and 6 (which are index 4 and 5)
xml_slides = prs.slides._sldIdLst
slides_list = list(xml_slides)
xml_slides.remove(slides_list[5]) # Remove original references
xml_slides.remove(slides_list[4]) # Remove original publications

prs.save(output_path)
print("Presentation generated successfully!")
